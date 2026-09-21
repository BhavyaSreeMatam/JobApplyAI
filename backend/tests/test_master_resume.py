"""Reading a resume as styled lines, and writing it back with different words.

The guarantee being tested is narrow and absolute: tailoring may change wording
and nothing else. So these check that the structure survives a round trip, that
facts of record are never offered for editing, and that an edit which would
repaginate the document can be taken back.
"""
import unittest
from pathlib import Path

from app.services.master_resume import (
    Line,
    MasterResume,
    Run,
    classify,
    parse_markup,
)


def line(index, text, **kwargs):
    runs = kwargs.pop("runs", None) or [Run(text)]
    return Line(index=index, runs=runs, **kwargs)


class MarkupTests(unittest.TestCase):
    """Inline weight has to survive the trip through the model."""

    def test_a_tagged_line_round_trips(self):
        original = [Run("Programming Languages: ", bold=True, size=10.0, font="Times"),
                    Run("Python, Java", size=10.0, font="Times")]
        rendered = Line(index=0, runs=original).markup()
        self.assertEqual(rendered, "<b>Programming Languages: </b>Python, Java")
        rebuilt = parse_markup(rendered, original)
        self.assertEqual([(r.text, r.bold) for r in rebuilt],
                         [("Programming Languages: ", True), ("Python, Java", False)])

    def test_an_edited_line_keeps_the_bold_label(self):
        original = [Run("Programming Languages: ", bold=True), Run("Python, Java")]
        rebuilt = parse_markup(
            "<b>Programming Languages: </b>Python, Java, C++, Spark", original
        )
        self.assertEqual(rebuilt[0].text, "Programming Languages: ")
        self.assertTrue(rebuilt[0].bold)
        self.assertEqual(rebuilt[1].text, "Python, Java, C++, Spark")
        self.assertFalse(rebuilt[1].bold)

    def test_new_runs_inherit_font_and_size_from_the_document(self):
        """Never a default - the size has to come from the resume itself."""
        original = [Run("x", bold=True, size=11.5, font="Georgia", color=(0.1, 0.2, 0.3))]
        rebuilt = parse_markup("<b>y</b> and more", original)
        self.assertEqual(rebuilt[0].size, 11.5)
        self.assertEqual(rebuilt[0].font, "Georgia")
        self.assertEqual(rebuilt[0].color, (0.1, 0.2, 0.3))

    def test_a_line_with_no_tags_becomes_one_run(self):
        rebuilt = parse_markup("plain text", [Run("original", size=9.0)])
        self.assertEqual(len(rebuilt), 1)
        self.assertEqual(rebuilt[0].text, "plain text")
        self.assertEqual(rebuilt[0].size, 9.0)

    def test_nested_tags_are_handled(self):
        rebuilt = parse_markup("<b><i>both</i></b>", [Run("x")])
        self.assertEqual([(r.text, r.bold, r.italic) for r in rebuilt if r.text],
                         [("both", True, True)])


class ClassificationTests(unittest.TestCase):
    """What may be re-worded, and what is a fact of record."""

    def build(self):
        lines = [
            line(0, "Alex Morgan", runs=[Run("Alex Morgan", bold=True, size=15.0)]),
            line(1, "alex@example.com || github.com/x"),
            line(2, "Graduate student in Computer Science at Example University with experience in XR."),
            line(3, "TECHNICAL SKILLS",
                 runs=[Run("TECHNICAL SKILLS", bold=True, underline=True)]),
            line(4, "Programming Languages: Python, Java",
                 runs=[Run("Programming Languages: ", bold=True), Run("Python, Java")]),
            line(5, "EXPERIENCE", runs=[Run("EXPERIENCE", bold=True, underline=True)]),
            line(6, "Research Assistant, Example University",
                 runs=[Run("Research Assistant, Example University", bold=True)],
                 right_text="2025 - Present"),
            line(7, "Built the cognitive-load pipeline in Python and Unity.", bullet="●"),
        ]
        classify(lines)
        return lines

    def test_kinds(self):
        kinds = [item.kind for item in self.build()]
        self.assertEqual(
            kinds,
            ["name", "contact", "summary", "heading", "skills", "heading", "entry", "bullet"],
        )

    def test_only_wording_lines_are_editable(self):
        editable = [item.index for item in self.build() if item.editable]
        self.assertEqual(editable, [2, 4, 7])

    def test_employers_titles_and_dates_are_never_offered_for_editing(self):
        entry = self.build()[6]
        self.assertEqual(entry.kind, "entry")
        self.assertFalse(entry.editable)
        self.assertEqual(entry.right_text, "2025 - Present")

    def test_section_and_entry_context_is_attached(self):
        bullet = self.build()[7]
        self.assertEqual(bullet.section, "EXPERIENCE")
        self.assertEqual(bullet.entry, "Research Assistant, Example University")


class ApplyTests(unittest.TestCase):
    def resume(self):
        lines = [
            line(0, "SKILLS", runs=[Run("SKILLS", bold=True, underline=True)]),
            line(1, "Python, Java"),
            line(2, "EXPERIENCE", runs=[Run("EXPERIENCE", bold=True, underline=True)]),
            line(3, "Engineer, Example Labs", runs=[Run("Engineer, Example Labs", bold=True)]),
            line(4, "Built retrieval pipelines in Python.", bullet="●"),
        ]
        classify(lines)
        return MasterResume("x.pdf", lines, {}, "pdf")

    def test_only_named_lines_change(self):
        resume = self.resume()
        changed = resume.apply({4: "Built RAG retrieval pipelines in Python."})
        self.assertEqual(changed, 1)
        self.assertEqual(resume.lines[1].text, "Python, Java")
        self.assertEqual(resume.lines[4].text, "Built RAG retrieval pipelines in Python.")

    def test_an_edit_to_a_protected_line_is_refused(self):
        """A model that tries to rename the employer is ignored, not obeyed."""
        resume = self.resume()
        self.assertEqual(resume.apply({3: "Senior Engineer, Example Laboratories"}), 0)
        self.assertEqual(resume.lines[3].text, "Engineer, Example Labs")
        self.assertEqual(resume.apply({0: "CORE SKILLS"}), 0)
        self.assertEqual(resume.lines[0].text, "SKILLS")

    def test_an_unknown_index_is_ignored(self):
        resume = self.resume()
        self.assertEqual(resume.apply({99: "nonsense"}), 0)

    def test_an_edit_can_be_taken_back_exactly(self):
        resume = self.resume()
        before = resume.lines[4].markup()
        resume.apply({4: "Built RAG retrieval pipelines in Python and PyTorch."})
        self.assertNotEqual(resume.lines[4].markup(), before)
        self.assertTrue(resume.revert(4))
        self.assertEqual(resume.lines[4].markup(), before)
        self.assertFalse(resume.revert(4))

    def test_growth_ranks_the_greediest_edit_first(self):
        """This is what the page-fit pass undoes when a resume overflows."""
        resume = self.resume()
        resume.apply({
            1: "Python, Java, Spark",
            4: "Built retrieval pipelines in Python, PyTorch and FAISS for vector search.",
        })
        grown = resume.growth()
        self.assertEqual(grown[0][0], 4)
        self.assertGreater(grown[0][1], grown[1][1])

    def test_blocks_carry_a_length_budget(self):
        blocks = {b["index"]: b for b in self.resume().blocks()}
        self.assertEqual(sorted(blocks), [1, 4])
        self.assertGreater(blocks[4]["max_chars"], len("Built retrieval pipelines in Python."))
        self.assertLess(blocks[4]["max_chars"], len("Built retrieval pipelines in Python.") * 1.4)

    def test_plain_text_includes_bullets_and_dates_for_scoring(self):
        text = self.resume().plain_text()
        self.assertIn("Built retrieval pipelines", text)
        self.assertIn("Engineer, Example Labs", text)


class DocxRoundTripTests(unittest.TestCase):
    """The format where the no-style-change promise is absolute."""

    def setUp(self):
        import tempfile

        import docx
        from docx.shared import Pt

        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "master.docx"

        document = docx.Document()

        def para(text, bold=False, underline=False, style=None, size=10):
            paragraph = document.add_paragraph(style=style)
            run = paragraph.add_run(text)
            run.bold, run.underline = bold, underline
            run.font.size = Pt(size)

        para("Test Person", bold=True, size=15)
        para("test@example.test || github.com/x")
        para("TECHNICAL SKILLS", bold=True, underline=True)
        mixed = document.add_paragraph()
        label = mixed.add_run("Programming Languages: ")
        label.bold = True
        label.font.size = Pt(10)
        mixed.add_run("Python, Java").font.size = Pt(10)
        para("EXPERIENCE", bold=True, underline=True)
        para("Engineer, Example Labs", bold=True)
        document.add_paragraph("Built retrieval pipelines in Python.", style="List Bullet")
        document.save(str(self.source))

    def parsed(self):
        from app.services.master_resume import load

        return load(self.source, "master.docx")

    def test_styles_are_read_from_the_document(self):
        lines = self.parsed().lines
        self.assertEqual(
            [line.kind for line in lines],
            ["name", "contact", "heading", "skills", "heading", "entry", "bullet"],
        )
        self.assertEqual(lines[3].markup(), "<b>Programming Languages: </b>Python, Java")
        self.assertEqual(lines[2].markup(), "<u><b>TECHNICAL SKILLS</b></u>")

    def test_a_style_applied_bullet_is_not_written_back_as_text(self):
        """Word numbers a list itself; re-adding the glyph shows two."""
        import docx

        resume = self.parsed()
        bullet = resume.lines[6]
        self.assertEqual(bullet.kind, "bullet")
        self.assertFalse(bullet.literal_bullet)

        resume.apply({6: "Built RAG retrieval pipelines in Python."})
        written = resume.write(Path(self.temp.name) / "out")
        text = docx.Document(str(written)).paragraphs[6].text
        self.assertEqual(text, "Built RAG retrieval pipelines in Python.")
        self.assertNotIn("•", text)

    def test_an_edit_keeps_every_run_property(self):
        import docx

        resume = self.parsed()
        resume.apply({3: "<b>Programming Languages: </b>Python, Java, C++, Unity"})
        written = resume.write(Path(self.temp.name) / "out")
        runs = docx.Document(str(written)).paragraphs[3].runs
        self.assertEqual(
            [(r.text, bool(r.bold), r.font.size.pt) for r in runs if r.text],
            [("Programming Languages: ", True, 10.0), ("Python, Java, C++, Unity", False, 10.0)],
        )

    def test_untouched_paragraphs_are_identical(self):
        import docx

        resume = self.parsed()
        resume.apply({6: "Built RAG retrieval pipelines in Python."})
        written = resume.write(Path(self.temp.name) / "out")
        before = docx.Document(str(self.source)).paragraphs
        after = docx.Document(str(written)).paragraphs
        for index in (0, 1, 2, 4, 5):
            with self.subTest(paragraph=index):
                self.assertEqual(before[index].text, after[index].text)
                self.assertEqual(
                    [(r.text, r.bold, r.underline) for r in before[index].runs],
                    [(r.text, r.bold, r.underline) for r in after[index].runs],
                )


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
