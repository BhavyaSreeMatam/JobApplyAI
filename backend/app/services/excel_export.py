from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.application import Application
from app.models.job import Job
from app.services.spreadsheet_safety import protect_text_cells


BACKEND_DIR = Path(__file__).resolve().parents[2]
EXPORT_DIRECTORY = BACKEND_DIR / "exports"
EXPORT_PATH = EXPORT_DIRECTORY / "job_applications.xlsx"


def format_datetime(value: datetime | None) -> str:
    if value is None:
        return ""

    return value.isoformat()


def format_status(value: object) -> str:
    return getattr(value, "value", str(value))


def style_worksheet(worksheet: Worksheet) -> None:
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        maximum_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        column_letter = column_cells[0].column_letter
        worksheet.column_dimensions[column_letter].width = min(
            maximum_length + 2,
            50,
        )


def export_applications(database: Session) -> Path:
    EXPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)

    statement = (
        select(Application, Job)
        .join(Job, Application.job_id == Job.id)
        .order_by(Application.created_at.desc())
    )

    records = database.execute(statement).all()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Applications"

    headers = [
        "Application ID",
        "Job ID",
        "Company",
        "Position",
        "Location",
        "Workplace Type",
        "Employment Type",
        "Salary Minimum",
        "Salary Maximum",
        "Currency",
        "Source",
        "Application URL",
        "Match Score",
        "Job Status",
        "Application Status",
        "Date Discovered",
        "Date Applied",
        "Resume Path",
        "Cover Letter Path",
        "Confirmation ID",
        "Failure Reason",
        "Last Updated",
    ]

    worksheet.append(headers)

    for application, job in records:
        worksheet.append(
            [
                application.id,
                job.id,
                job.company,
                job.title,
                job.location or "",
                job.workplace_type or "",
                job.employment_type or "",
                job.salary_min,
                job.salary_max,
                job.salary_currency or "",
                job.source,
                job.application_url,
                job.match_score,
                format_status(job.status),
                format_status(application.status),
                format_datetime(job.discovered_at),
                format_datetime(application.applied_at),
                application.resume_path or "",
                application.cover_letter_path or "",
                application.confirmation_id or "",
                application.failure_reason or "",
                format_datetime(application.updated_at),
            ]
        )

    style_worksheet(worksheet)
    protect_text_cells(workbook)

    temporary_path = EXPORT_DIRECTORY / "job_applications.tmp.xlsx"
    workbook.save(temporary_path)
    temporary_path.replace(EXPORT_PATH)

    return EXPORT_PATH
